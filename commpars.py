import os
import re
import time
from urllib.parse import urlparse

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


load_dotenv()

API_KEY = os.getenv("YOUTUBE_API_KEY")
BASE_URL = "https://www.googleapis.com/youtube/v3"

if not API_KEY:
    raise RuntimeError("Не найден YOUTUBE_API_KEY в .env")


def api_get(endpoint: str, params: dict) -> dict:
    """
    Универсальный GET-запрос к YouTube Data API.
    """
    params = params.copy()
    params["key"] = API_KEY

    response = requests.get(f"{BASE_URL}/{endpoint}", params=params, timeout=30)

    if response.status_code != 200:
        raise RuntimeError(
            f"Ошибка API {response.status_code}:\n{response.text}"
        )

    return response.json()


def extract_channel_ref(text: str) -> str:
    """
    Принимает:
    - @handle
    - https://www.youtube.com/@handle
    - https://www.youtube.com/channel/UC...
    - UC...
    """
    text = text.strip()

    if text.startswith("http"):
        parsed = urlparse(text)
        path = parsed.path.strip("/")

        if path.startswith("@"):
            return path

        if path.startswith("channel/"):
            return path.split("/")[1]

        if path.startswith("user/"):
            return path.split("/")[1]

        if path.startswith("c/"):
            return path.split("/")[1]

    return text


def get_channel_info(channel_input: str) -> dict:
    """
    Получает channel_id, название канала и uploads playlist.
    """
    channel_ref = extract_channel_ref(channel_input)

    attempts = []

    if channel_ref.startswith("UC"):
        attempts.append({"id": channel_ref})
    elif channel_ref.startswith("@"):
        attempts.append({"forHandle": channel_ref})
        attempts.append({"forHandle": channel_ref.lstrip("@")})
    else:
        attempts.append({"forHandle": channel_ref})
        attempts.append({"forHandle": f"@{channel_ref}"})
        attempts.append({"forUsername": channel_ref})

    last_data = None

    for attempt in attempts:
        params = {
            "part": "snippet,contentDetails",
            "maxResults": 1,
        }
        params.update(attempt)

        data = api_get("channels", params)
        last_data = data

        items = data.get("items", [])
        if items:
            item = items[0]

            return {
                "channel_id": item["id"],
                "channel_title": item["snippet"]["title"],
                "uploads_playlist_id": item["contentDetails"]["relatedPlaylists"]["uploads"],
            }

    raise RuntimeError(
        "Канал не найден. Лучше вставь ссылку вида:\n"
        "https://www.youtube.com/@handle\n"
        "или:\n"
        "https://www.youtube.com/channel/UC..."
    )


def get_all_videos_from_uploads_playlist(playlist_id: str) -> list[dict]:
    """
    Получает все видео из uploads-плейлиста канала.
    """
    videos = []
    page_token = None

    while True:
        params = {
            "part": "snippet,contentDetails",
            "playlistId": playlist_id,
            "maxResults": 50,
        }

        if page_token:
            params["pageToken"] = page_token

        data = api_get("playlistItems", params)

        for item in data.get("items", []):
            snippet = item.get("snippet", {})
            resource_id = snippet.get("resourceId", {})
            video_id = resource_id.get("videoId")

            if not video_id:
                continue

            title = snippet.get("title", "")

            # Пропускаем удалённые/приватные ролики
            if title in ["Deleted video", "Private video"]:
                continue

            videos.append({
                "video_id": video_id,
                "video_title": title,
                "published_at": snippet.get("publishedAt", ""),
                "video_url": f"https://www.youtube.com/watch?v={video_id}",
            })

        page_token = data.get("nextPageToken")

        print(f"Собрано видео: {len(videos)}")

        if not page_token:
            break

        time.sleep(0.1)

    return videos


def get_replies_for_comment(parent_comment_id: str, video: dict) -> list[dict]:
    """
    Получает все ответы на один верхнеуровневый комментарий.
    """
    replies = []
    page_token = None

    while True:
        params = {
            "part": "snippet",
            "parentId": parent_comment_id,
            "maxResults": 100,
            "textFormat": "plainText",
        }

        if page_token:
            params["pageToken"] = page_token

        try:
            data = api_get("comments", params)
        except RuntimeError as e:
            print(f"Не смог получить ответы для комментария {parent_comment_id}: {e}")
            return replies

        for item in data.get("items", []):
            snippet = item.get("snippet", {})

            replies.append({
                "video_id": video["video_id"],
                "video_title": video["video_title"],
                "video_url": video["video_url"],
                "comment_id": item.get("id", ""),
                "parent_comment_id": parent_comment_id,
                "is_reply": True,
                "author": snippet.get("authorDisplayName", ""),
                "author_channel_url": snippet.get("authorChannelUrl", ""),
                "text": snippet.get("textDisplay", ""),
                "like_count": snippet.get("likeCount", 0),
                "published_at": snippet.get("publishedAt", ""),
                "updated_at": snippet.get("updatedAt", ""),
                "error": "",
            })

        page_token = data.get("nextPageToken")

        if not page_token:
            break

        time.sleep(0.1)

    return replies


def get_comments_for_video(video: dict) -> tuple[list[dict], list[dict]]:
    """
    Получает все комментарии конкретного видео:
    - верхнеуровневые комментарии
    - ответы на них

    Возвращает:
    comments, errors
    """
    comments = []
    errors = []
    page_token = None

    while True:
        params = {
            "part": "snippet",
            "videoId": video["video_id"],
            "maxResults": 100,
            "order": "time",
            "textFormat": "plainText",
        }

        if page_token:
            params["pageToken"] = page_token

        try:
            data = api_get("commentThreads", params)
        except RuntimeError as e:
            error_text = str(e)

            if "commentsDisabled" in error_text:
                print(f"Комментарии отключены: {video['video_title']}")
                errors.append({
                    "video_id": video["video_id"],
                    "video_title": video["video_title"],
                    "video_url": video["video_url"],
                    "error": "commentsDisabled",
                })
                return comments, errors

            print(f"Ошибка при получении комментариев к видео {video['video_id']}: {e}")

            errors.append({
                "video_id": video["video_id"],
                "video_title": video["video_title"],
                "video_url": video["video_url"],
                "error": error_text[:1000],
            })

            return comments, errors

        for item in data.get("items", []):
            thread_snippet = item.get("snippet", {})
            top_comment = thread_snippet.get("topLevelComment", {})
            top_snippet = top_comment.get("snippet", {})

            top_comment_id = top_comment.get("id", "")

            comments.append({
                "video_id": video["video_id"],
                "video_title": video["video_title"],
                "video_url": video["video_url"],
                "comment_id": top_comment_id,
                "parent_comment_id": "",
                "is_reply": False,
                "author": top_snippet.get("authorDisplayName", ""),
                "author_channel_url": top_snippet.get("authorChannelUrl", ""),
                "text": top_snippet.get("textDisplay", ""),
                "like_count": top_snippet.get("likeCount", 0),
                "published_at": top_snippet.get("publishedAt", ""),
                "updated_at": top_snippet.get("updatedAt", ""),
                "error": "",
            })

            total_reply_count = thread_snippet.get("totalReplyCount", 0)

            if total_reply_count > 0 and top_comment_id:
                replies = get_replies_for_comment(top_comment_id, video)
                comments.extend(replies)

        page_token = data.get("nextPageToken")

        if not page_token:
            break

        time.sleep(0.1)

    return comments, errors


def safe_filename(name: str) -> str:
    """
    Делает безопасное имя файла.
    """
    name = re.sub(r'[\\/*?:"<>|]', "_", name)
    name = name.strip()
    return name[:100] or "youtube_comments"


def save_to_xlsx(rows: list[dict], errors: list[dict], filename: str):
    """
    Сохраняет комментарии и ошибки в XLSX.
    """
    wb = Workbook()

    ws = wb.active
    ws.title = "comments"

    headers = [
        "video_id",
        "video_title",
        "video_url",
        "comment_id",
        "parent_comment_id",
        "is_reply",
        "author",
        "author_channel_url",
        "text",
        "like_count",
        "published_at",
        "updated_at",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))

        ws.column_dimensions[column_letter].width = min(max_length + 2, 70)

    # Лист с ошибками
    ws_errors = wb.create_sheet("errors")

    error_headers = [
        "video_id",
        "video_title",
        "video_url",
        "error",
    ]

    ws_errors.append(error_headers)

    for cell in ws_errors[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for error in errors:
        ws_errors.append([
            error.get("video_id", ""),
            error.get("video_title", ""),
            error.get("video_url", ""),
            error.get("error", ""),
        ])

    ws_errors.freeze_panes = "A2"
    ws_errors.auto_filter.ref = ws_errors.dimensions

    for row in ws_errors.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_cells in ws_errors.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))

        ws_errors.column_dimensions[column_letter].width = min(max_length + 2, 70)

    wb.save(filename)


def parse_channel_to_xlsx(channel_input: str, progress_callback=None) -> dict:
    """
    Запускает полный парсинг канала и сохраняет результат в XLSX.
    Может отправлять прогресс через progress_callback.
    """

    channel_info = get_channel_info(channel_input)

    if progress_callback:
        progress_callback({
            "stage": "channel_found",
            "channel_title": channel_info["channel_title"],
            "videos_count": 0,
            "comments_count": 0,
            "errors_count": 0,
            "current_video_index": 0,
            "current_video_title": "",
        })

    videos = get_all_videos_from_uploads_playlist(
        channel_info["uploads_playlist_id"]
    )

    all_comments = []
    all_errors = []

    if progress_callback:
        progress_callback({
            "stage": "videos_loaded",
            "channel_title": channel_info["channel_title"],
            "videos_count": len(videos),
            "comments_count": 0,
            "errors_count": 0,
            "current_video_index": 0,
            "current_video_title": "",
        })

    for index, video in enumerate(videos, start=1):
        print(f"[{index}/{len(videos)}] {video['video_title']}")

        if progress_callback:
            progress_callback({
                "stage": "video_started",
                "channel_title": channel_info["channel_title"],
                "videos_count": len(videos),
                "comments_count": len(all_comments),
                "errors_count": len(all_errors),
                "current_video_index": index,
                "current_video_title": video["video_title"],
            })

        video_comments, video_errors = get_comments_for_video(video)

        all_comments.extend(video_comments)
        all_errors.extend(video_errors)

        print(
            f"Комментарии у видео: {len(video_comments)} | "
            f"Всего собрано: {len(all_comments)} | "
            f"Ошибок: {len(all_errors)}"
        )

        if progress_callback:
            progress_callback({
                "stage": "video_finished",
                "channel_title": channel_info["channel_title"],
                "videos_count": len(videos),
                "comments_count": len(all_comments),
                "errors_count": len(all_errors),
                "current_video_index": index,
                "current_video_title": video["video_title"],
            })

    filename = safe_filename(channel_info["channel_title"]) + "_comments.xlsx"

    if progress_callback:
        progress_callback({
            "stage": "saving_xlsx",
            "channel_title": channel_info["channel_title"],
            "videos_count": len(videos),
            "comments_count": len(all_comments),
            "errors_count": len(all_errors),
            "current_video_index": len(videos),
            "current_video_title": "",
        })

    save_to_xlsx(all_comments, all_errors, filename)

    if progress_callback:
        progress_callback({
            "stage": "xlsx_saved",
            "channel_title": channel_info["channel_title"],
            "videos_count": len(videos),
            "comments_count": len(all_comments),
            "errors_count": len(all_errors),
            "current_video_index": len(videos),
            "current_video_title": "",
        })

    return {
        "filename": filename,
        "channel_title": channel_info["channel_title"],
        "videos_count": len(videos),
        "comments_count": len(all_comments),
        "errors_count": len(all_errors),
    }


def main():
    channel_input = input("Вставь @handle, ссылку на канал или channel_id: ").strip()

    result = parse_channel_to_xlsx(channel_input)

    print("\nГотово.")
    print(f"Канал: {result['channel_title']}")
    print(f"Видео: {result['videos_count']}")
    print(f"Комментариев: {result['comments_count']}")
    print(f"Ошибок: {result['errors_count']}")
    print(f"Файл: {result['filename']}")


if __name__ == "__main__":
    main()